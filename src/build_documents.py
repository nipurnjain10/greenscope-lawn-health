"""Build the report, presentation, and submission copies from the saved experiment."""

from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.util import Emu, Inches as PptInches, Pt as PptPt

from src.config import CLASS_ORDER, DATA_DIR, FIGURE_DIR, REPORT_DIR, ROOT, SUBMISSION_DIR
from src.utilities import read_json

INK = RGBColor(0x1C, 0x28, 0x33)
GREEN = RGBColor(0x1E, 0x4D, 0x36)


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _load() -> dict:
    metrics = read_json(REPORT_DIR / "model_metrics.json")
    data_meta = read_json(DATA_DIR / "metadata.json")
    run_meta = read_json(ROOT / "models" / "run_metadata.json")
    group = read_json(ROOT / "group_info.json")
    importance = None
    importance_path = REPORT_DIR / "permutation_importance.csv"
    if importance_path.exists():
        import pandas as pd

        importance = pd.read_csv(importance_path)
    return {
        "metrics": metrics,
        "data": data_meta,
        "run": run_meta,
        "group": group,
        "importance": importance,
    }


class Blocks:
    def __init__(self):
        self.items: list[tuple] = []

    def heading(self, level: int, text: str):
        self.items.append(("h", level, text))

    def paragraph(self, text: str):
        self.items.append(("p", text))

    def bullets(self, rows: list[str]):
        self.items.append(("b", rows))

    def table(self, headers: list[str], rows: list[list[str]]):
        self.items.append(("t", headers, rows))

    def image(self, path: Path, caption: str):
        self.items.append(("i", path, caption))

    def markdown(self) -> str:
        lines: list[str] = []
        for item in self.items:
            kind = item[0]
            if kind == "h":
                level, text = item[1], item[2]
                lines.extend(["", "#" * level + " " + text, ""])
            elif kind == "p":
                lines.extend([item[1], ""])
            elif kind == "b":
                lines.extend(f"- {row}" for row in item[1])
                lines.append("")
            elif kind == "t":
                headers, rows = item[1], item[2]
                lines.append("| " + " | ".join(headers) + " |")
                lines.append("| " + " | ".join("---" for _ in headers) + " |")
                for row in rows:
                    lines.append("| " + " | ".join(row) + " |")
                lines.append("")
            elif kind == "i":
                path, caption = item[1], item[2]
                if path.exists():
                    relative = path.relative_to(ROOT).as_posix()
                    lines.append(f"![{caption}]({relative})")
                    lines.append("")
                    lines.append(f"*{caption}*")
                    lines.append("")
        return "\n".join(lines).strip() + "\n"

    def docx(self, path: Path):
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = INK
        for item in self.items:
            kind = item[0]
            if kind == "h":
                heading = document.add_heading(item[2], level=item[1])
                for run in heading.runs:
                    run.font.color.rgb = GREEN
            elif kind == "p":
                document.add_paragraph(item[1])
            elif kind == "b":
                for row in item[1]:
                    document.add_paragraph(row, style="List Bullet")
            elif kind == "t":
                headers, rows = item[1], item[2]
                table = document.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = "Table Grid"
                for index, header in enumerate(headers):
                    table.rows[0].cells[index].text = header
                for row_index, row in enumerate(rows, start=1):
                    for column, value in enumerate(row):
                        table.rows[row_index].cells[column].text = value
                document.add_paragraph("")
            elif kind == "i" and item[1].exists():
                document.add_picture(str(item[1]), width=Inches(6.2))
                caption = document.add_paragraph(item[2])
                caption.italic = True
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)


def _metric_rows(metrics: dict, split: str) -> list[list[str]]:
    rows = []
    for name in [
        "logistic_regression",
        "random_forest",
        "gradient_boosting",
        "svm_rbf",
        "soft_voting",
    ]:
        block = metrics["models"][name][split]
        rows.append(
            [
                name.replace("_", " "),
                _pct(block["accuracy"]),
                _pct(block["macro_precision"]),
                _pct(block["macro_recall"]),
                _pct(block["macro_f1"]),
                _pct(block["weighted_f1"]),
                str(block["high_risk_predicted_low"]),
            ]
        )
    return rows


def build_report(bundle: dict) -> Blocks:
    metrics = bundle["metrics"]
    data = bundle["data"]
    run = bundle["run"]
    group = bundle["group"]
    soft = metrics["models"]["soft_voting"]["test"]
    counts = data["class_counts_main"]
    doc = Blocks()
    doc.heading(1, "Using Machine Learning to Improve the Business of a Chain of Gardening Works Like The Grounds Guys")
    doc.paragraph("Technical title: An Explainable Ensemble Learning Framework for Plant Health Risk Prediction and Cost-Aware Landscaping Decisions.")
    doc.paragraph("Application: GreenScope — Lawn Health & Service Planning. An independent academic prototype. It is not software of The Grounds Guys and is not affiliated with that company.")
    doc.paragraph(f"Group number: {group['group_number']}")
    doc.paragraph(f"Course: {group['course']}, {group['semester']}")
    doc.paragraph(f"Institute: {group['institute']}")
    doc.paragraph(f"Instructor: {group['instructor']}")
    doc.paragraph(f"Submission date: {group['submission_date']}")
    doc.paragraph("Document generated from the saved experiment files. Replace the bracketed group fields in group_info.json and rebuild this document before submission.")
    doc.heading(1, "Individual contributions")
    doc.paragraph(
        "The professor requires each member's name, a one- or two-sentence contribution, and the pages where that work appears. "
        "The rows below are a proposed division of work, not a record of work already done by a named person. "
        "Replace each contribution sentence after the group confirms it. Printed page numbers depend on the word processor, so the section locations must be checked on the paginated copy."
    )
    doc.table(
        ["Name", "Enrollment", "Proposed role", "Contribution to confirm", "Where to confirm the pages"],
        [
            [member["name"], member["enrollment"], member["proposed_role"], member["contribution"], member["report_location"]]
            for member in group["members"]
        ],
    )
    doc.paragraph("All five members are expected to review the interpretation, the report, and the presentation. Do not leave the placeholder sentences in the submitted copy.")

    doc.heading(1, "1. Abstract")
    doc.paragraph(
        "This semester project builds GreenScope, a local decision-support prototype for lawn inspection. "
        "A reproducible generator creates 10,000 synthetic inspections and a separate 2,000-row robustness file. "
        "Four supervised classifiers — logistic regression, a random forest, gradient boosting, and a calibrated RBF support-vector machine — predict one simulated risk class: Low, Moderate, or High. "
        "Soft voting averages their probabilities. The application explains the model with permutation importance and a local baseline contrast, then applies transparent water, nutrient, and drainage screens and an editable cost sheet."
    )
    doc.paragraph(
        f"On the untouched test set of {soft['support_total']} rows, soft voting accuracy is {_pct(soft['accuracy'])} and macro-F1 is {_pct(soft['macro_f1'])}. "
        + metrics["narratives"]["performance"]
    )
    doc.paragraph("These figures describe the synthetic benchmark only. The system has not been validated on customer lawns and no profit claim is made.")

    doc.heading(1, "2. Problem statement and motivation")
    doc.paragraph(
        "A landscaping crew can record area, grass type, light, a soil test, recent weather, maintenance, and what the lawn looks like. "
        "Turning that sheet into a priority and a budget is still a judgement. The academic question is whether a small ensemble can make that judgement explicit and checkable on a benchmark whose assumptions are written down."
    )
    doc.paragraph(
        "The original course title refers to the business problems of a lawn-care chain. The software identity is GreenScope. "
        "The company name is context for the problem, not a claim of access to its customers or its accounts."
    )

    doc.heading(1, "3. Business context and project scope")
    doc.paragraph(
        "The prototype supports two jobs. A field user records an inspection and receives a risk class, the four model votes, an explanation, screening notes, actions, and an illustrative cost. "
        "A reader of the research page sees the same models evaluated on one held-out test set, plus a shifted robustness file, an ablation, calibration, and error examples."
    )
    doc.paragraph("Out of scope: accounts, authentication, live weather, chemical prescriptions, pesticide advice, and any statement that the model raises revenue.")

    doc.heading(1, "4. Research objectives and questions")
    doc.paragraph("Central question: can an ensemble of diverse supervised classifiers produce a useful and explainable lawn-health risk assessment on a source-informed synthetic benchmark, and how can that assessment support transparent, cost-aware decisions?")
    doc.bullets(
        [
            "Does soft voting outperform the best individual model on the held-out synthetic test set?",
            "Which recorded features does the model depend on?",
            "Does performance change when environmental draws are shifted?",
            "How can the assessment become a maintenance priority and an illustrative budget without pretending to be a prescription?",
        ]
    )

    doc.heading(1, "5. Novelty and intended contribution")
    doc.paragraph(
        "The contribution is the combination of a documented synthetic lawn process, four different classifiers on one target and one split, "
        "equal-weight voting that is allowed to lose, model explanations that are not called causes, rule-based stress screens that are not called classifiers, "
        "and a cost sheet whose prices are editable and labelled as illustrations. The novelty is not a claimed accuracy and not a commercial platform."
    )

    doc.heading(1, "6. Related work and source-grounded background")
    doc.paragraph(
        "University of Maryland Extension describes site evaluation, about six or more hours of direct sun, slow drainage, compaction, and a favourable pH band near 6.0–6.8, with lime discussed below 5.5. "
        "The same service describes about 25 mm of water per week for actively growing turf, drought signs, core aeration, and thatch beyond about one-half inch. "
        "Rutgers Cooperative Extension fact sheets FS797, FS633, and FS102 describe soil testing, the absence of a routine soil test for turfgrass nitrogen need, soil-test-based phosphorus and potassium decisions, mowing, and weed identification. "
        "The FAO EcoCrop home page lists environmental factors used to match crops. No EcoCrop species table was downloaded. Full notes are in data/source_register.md."
    )

    doc.heading(1, "7. Dataset sources and provenance")
    doc.paragraph(data["provenance"])
    doc.paragraph(
        f"Generator {data['generator_version']}, primary seed {data['seed_main']}, robustness seed {data['seed_robustness']}. "
        f"Primary fingerprint {data['fingerprint_sha256_main'][:12]}… and robustness fingerprint {data['fingerprint_sha256_robustness'][:12]}… "
        "No external dataset row is in either CSV. Candidate public turf files were identified and not downloaded; the reasons are in the source register."
    )

    doc.heading(1, "8. Synthetic data-generation methodology")
    doc.paragraph(
        "Hidden normal variables create dependence among heat, water supply, soil chemistry, care, and site. "
        "Observed rainfall, moisture, irrigation, nutrients, compaction, and drainage are noisy functions of those variables, so the columns are not independent. "
        "The label is not a function of the hidden variables. Generator 1.0.0 did that and a training-only pilot could not recover the label (logistic regression about 0.64 validation accuracy; a deep forest memorised the training rows). "
        "Generator 1.1.0 scores the measurements: dryness together with heat, pH distance from a mildly acidic centre together with low phosphorus, and poor drainage together with wet soil, plus shade, neglected mowing or thatch, a noisy symptom term, and Gaussian noise with standard deviation "
        f"{data['label_noise_sd']}."
    )
    doc.paragraph(data["class_rule"] + f" Frozen cuts: Low below {data['class_thresholds']['low_lt']:.3f}, High at or above {data['class_thresholds']['high_ge']:.3f}.")
    doc.paragraph("Primary class counts: " + ", ".join(f"{label} {counts.get(label, 0)}" for label in CLASS_ORDER) + f", from {data['n_main']} rows.")
    doc.paragraph(
        "Robustness shifts: " + " ".join(data["robustness_shifts"])
        + " Robustness class counts: "
        + ", ".join(f"{label} {data['class_counts_robustness'].get(label, 0)}" for label in CLASS_ORDER)
        + "."
    )
    doc.image(FIGURE_DIR / "fig_class_distribution.png", "Figure 1. Simulated class counts in the primary file. The near-balance is produced by the tertile rule.")
    doc.paragraph("Continuous draws are clipped to the documented bounds inside the generator. Validation does not clip rows afterwards to force a pass.")

    doc.heading(1, "9. Data dictionary and preprocessing")
    doc.paragraph(
        "The model sees 23 inputs in five groups: lawn and site, soil measurements, environment, maintenance, and visible observations. "
        "Units, bounds, and categories are in data/data_dictionary.md. "
        "Pipelines impute numeric medians and categorical modes, one-hot encode categories with unknown levels ignored, and scale numerics only for logistic regression and the SVM. "
        "Imputers and encoders are fit inside the training partition. The target and the intermediate score used to assign it are not features."
    )
    doc.paragraph(
        "Laboratory pH and nutrient values are missing more often when the last service was at least 100 days earlier. "
        f"Primary missingness: {data['missingness_main']}. "
        "An unknown value in the application is passed through as missing. It is not replaced with a healthy constant before the pipeline."
    )

    doc.heading(1, "10. Experimental setup and split")
    doc.paragraph(metrics["narratives"]["split"])
    doc.paragraph(metrics["evaluation_protocol"])
    doc.paragraph(
        "Software recorded with the run: Python "
        + run["python_version"]
        + "; "
        + ", ".join(f"{name} {version}" for name, version in run["package_versions"].items())
        + f". Training timestamp {run['created_at_utc']}."
    )

    doc.heading(1, "11. Four classifier descriptions")
    for name, text in run["hyperparameters"].items():
        if name == "voting":
            continue
        doc.paragraph(f"{name.replace('_', ' ').capitalize()}: {text}.")
    doc.paragraph(
        "Logistic regression tests a linear boundary. The forest and the boosting model can use interactions such as heat with dryness. "
        "The SVM uses a nonlinear RBF kernel. Its probabilities come from a three-fold sigmoid calibrator fit on the training rows, not from the test set. "
        "A validation-only grid preferred gradient boosting with 300 trees, depth 3, and learning rate 0.05 over a shorter schedule. Deeper forests did not move validation accuracy much past the mid-0.70s, so the reported forest stays depth-limited."
    )

    doc.heading(1, "12. Soft voting")
    doc.paragraph(run["hyperparameters"]["voting"])
    doc.paragraph(
        "Soft voting is the principal ensemble: for each class, average the four probabilities and choose the largest. "
        "Soft voting is the only ensemble. Removing one model for the ablation does not retrain the other three."
    )

    doc.heading(1, "13. Metrics")
    doc.paragraph(
        "Accuracy is the share of test rows whose predicted class matches the simulated label. "
        "Precision, recall, and F1 are reported per class and as macro and weighted averages. "
        "Macro-F1 treats the three classes equally. A confusion matrix shows which simulated class was confused with which prediction. "
        "The high-risk row predicted as low risk is the miss emphasised in the discussion. "
        "Bootstrap intervals resample the stored predictions and do not refit the models. They are not real-world confidence intervals."
    )

    doc.heading(1, "14. Results and model comparison")
    doc.table(
        ["Model", "Accuracy", "Macro-precision", "Macro-recall", "Macro-F1", "Weighted-F1", "High called low"],
        _metric_rows(metrics, "test"),
    )
    doc.image(FIGURE_DIR / "fig_model_comparison.png", "Figure 2. Test-set scores for the four classifiers and soft voting. All bars use the same held-out rows.")
    doc.paragraph(metrics["narratives"]["performance"])
    doc.paragraph(
        "Validation scores from the training-only fit are stored in model_metrics.json under validation_before_refit. "
        "They are not plotted against the test scores. The table above is the test set after the refit."
    )

    doc.heading(1, "15. Confusion matrices and class-wise performance")
    for name in ["logistic_regression", "random_forest", "gradient_boosting", "svm_rbf", "soft_voting"]:
        doc.image(
            FIGURE_DIR / f"fig_confusion_{name}.png",
            f"Figure. Confusion matrix for {name.replace('_', ' ')} on the test set. Rows are the simulated class and columns are the prediction.",
        )
    doc.image(FIGURE_DIR / "fig_per_class_soft_voting.png", "Figure. Soft-voting precision, recall, and F1 by simulated class.")
    doc.image(FIGURE_DIR / "fig_recall_by_class.png", "Figure. Test recall by class for every reported model.")
    doc.paragraph(metrics["narratives"]["per_class"])

    doc.heading(1, "16. Explainability")
    doc.image(FIGURE_DIR / "fig_feature_importance.png", "Figure. Permutation importance for soft voting on the test set. Error bars are the standard deviation across repeats.")
    doc.paragraph(metrics["narratives"]["importance"])
    doc.paragraph(
        "For a single inspection the application replaces one observed value with the fitting-set median or mode and records the change in the soft-voting probability of the predicted class. "
        "A missing measurement is not given a contrast. The explanation is dependence of this model, not a cause of the lawn's condition."
    )

    doc.heading(1, "17. Robustness")
    doc.table(
        ["Model", "Accuracy", "Macro-precision", "Macro-recall", "Macro-F1", "Weighted-F1", "High called low"],
        _metric_rows(metrics, "robustness"),
    )
    doc.image(FIGURE_DIR / "fig_robustness.png", "Figure. Accuracy on the standard test set and on the separately generated robustness file.")
    doc.paragraph(metrics["narratives"]["robustness"])

    doc.heading(1, "18. Ablation and error analysis")
    doc.image(FIGURE_DIR / "fig_ablation.png", "Figure. Test macro-F1 for the full soft vote and for each leave-one-out average. The remaining models are not retrained.")
    doc.paragraph(metrics["narratives"]["ablation"])
    doc.image(FIGURE_DIR / "fig_disagreement.png", "Figure. Agreement among the four base models on the test set.")
    disagreement = metrics["disagreement"]
    doc.paragraph(
        f"Of {disagreement['n']} test rows, {disagreement['unanimous']} are unanimous, "
        f"{disagreement['majority_with_dissent']} have a single dissent, and {disagreement['split']} have no three-vote majority."
    )
    doc.paragraph(
        "Saved error examples include high-risk rows called low, rows where the models disagree, and high-risk rows the ensemble got right. "
        "They are in reports/error_examples.json. Recurring difficulty is expected near the score cutoffs, where the noise term can move a row across a boundary, and where a laboratory value was hidden."
    )

    doc.heading(1, "19. Calibration")
    doc.image(FIGURE_DIR / "fig_calibration_soft_voting.png", "Figure. One-versus-rest reliability of the soft-voting probabilities on the test set.")
    doc.paragraph(metrics["narratives"]["calibration"])

    doc.heading(1, "20. Decision support and recommendations")
    doc.paragraph(
        "Water, nutrient, and drainage screens are written rules. They do not emit probabilities and they were not fit as classifiers. "
        "Water rules use low moisture with heat, a dry week with little irrigation, a high wilting score, and frequent irrigation that has not raised moisture. "
        "Nutrient rules use the Maryland pH discussion and project-only mg/kg bands that are not laboratory sufficiency classes. "
        "Rutgers' statement that a routine soil test does not establish nitrogen need is shown whenever a low nitrogen value is discussed. "
        "Drainage rules use the worker drainage score, compaction, fine texture, and wet soil. "
        "Actions ask for a soil test, an irrigation inspection, a drainage look, aeration when compaction is high, a sunlight review, a thatch check above about 12.7 mm, and a follow-up. "
        "No action states a fertiliser, lime, sulfur, or pesticide rate. A high overall class does not by itself create a nutrient action."
    )
    doc.paragraph(
        "Priority is urgency weight times evidence strength. Cost is shown separately and does not change the rank. "
        "The follow-up interval is 14, 28, or 56 days according to the class and the screens. It is an operational heuristic. "
        "Comparing two saved assessments shows input changes and a new prediction. It does not estimate a treatment effect."
    )

    doc.heading(1, "21. Business cost-estimation methodology")
    doc.paragraph(
        "Each priced action is a flat amount plus a per-square-metre amount times the entered lawn area. "
        "The default books are labelled illustrative Indian rupees and a separate illustrative US dollar schedule. "
        "The dollar schedule is not an exchange-rate conversion. "
        "A minimal plan prices the high-urgency actions. A comprehensive plan adds the medium-urgency actions. "
        "Users can edit every flat fee and every per-square-metre rate. "
        "The result is a calculated illustration. It is not an invoice, a replacement-cost study, or evidence of savings."
    )

    doc.heading(1, "22. Application design and demonstration")
    doc.paragraph(
        "The Streamlit application has an overview, a lawn assessment with recommendations and costs, a demonstration lab, and a combined results and data page. "
        "The demonstration loads one held-out test row, runs the four saved pipelines, averages the probabilities, and only then shows the simulated benchmark label. "
        "The operational page does not show that label, because a new lawn does not have one. "
        "Models are loaded from disk. Page changes do not refit them. "
        "A fingerprint check warns if the CSV and the saved run diverge."
    )

    doc.heading(1, "23. Limitations and threats to validity")
    doc.bullets(
        [
            "The labels are simulated. Agreement with them is not agronomic truth.",
            "Extension pages are from Maryland, New Jersey, and an FAO index. They are not a sample of Indian lawns, and the rupee prices are not a local survey.",
            "Project mg/kg bands are not a soil-test method.",
            "Ordinal scores are treated as equally spaced numbers by the linear model.",
            "Median imputation fills missing laboratory values with a training statistic. That statistic was not measured on the lawn.",
            "Probability estimates can be misread as real-world confidence. The calibration plot is the check, and it uses synthetic labels.",
            "The robustness file changes several things at once, so a drop cannot be attributed to one factor.",
            "Ablation does not retrain the remaining models.",
            "Bootstrap intervals do not refit the models and do not describe customer lawns.",
            "Validation was used to revise the generator once and to choose the boosting schedule. The test set was not used for those choices, but the validation split of the final generator was inspected, so the validation metrics are not a second untouched test.",
        ]
    )

    doc.heading(1, "24. Conclusion and recommendations")
    doc.paragraph(metrics["narratives"]["performance"])
    doc.paragraph(metrics["narratives"]["ablation"])
    doc.paragraph(metrics["narratives"]["robustness"])
    doc.paragraph(
        "Recommendations for anyone using the prototype: treat a disagreement among the four models as a reason to reinspect; "
        "do not convert the class into a fertiliser rate; obtain a laboratory soil test before a nutrient or lime decision; "
        "read a high-to-low error as the operational miss that matters on this benchmark; "
        "and edit the price book before discussing a budget. "
        "Recommendations for the work itself: the benchmark is strong enough to compare models and weak as evidence about real lawns. "
        "A later study needs real inspections with a documented labelling protocol before anyone claims field accuracy or a business result."
    )

    doc.heading(1, "References")
    doc.bullets(
        [
            "University of Maryland Extension. Starting a New Lawn. https://extension.umd.edu/resource/starting-new-lawn. Accessed 22 October 2026.",
            "University of Maryland Extension. Maintaining an Established Lawn. https://extension.umd.edu/resource/maintaining-established-lawn. Accessed 22 October 2026.",
            "Heckman, J. Soil Testing for Home Lawns and Gardens (FS797). Rutgers Cooperative Extension. https://njaes.rutgers.edu/fs797/. Accessed 22 October 2026.",
            "Heckman, J. and Murphy, J. A. Fertilizing the Home Lawn (FS633). Rutgers Cooperative Extension. https://njaes.rutgers.edu/FS633/. Accessed 22 October 2026.",
            "Murphy, J. A. Your Lawn and Its Care (FS102). Rutgers Cooperative Extension. https://njaes.rutgers.edu/fs102/. Accessed 22 October 2026.",
            "Food and Agriculture Organization of the United Nations. EcoCrop home page. https://ecocrop.apps.fao.org/ecocrop/srv/en/home. Accessed 22 October 2026. No species records were downloaded.",
        ]
    )
    return doc


def build_experiment_markdown(bundle: dict) -> str:
    metrics = bundle["metrics"]
    lines = [
        "# Experiment analysis",
        "",
        "This note is generated from `reports/model_metrics.json`. It is the same evaluation the research page reads.",
        "",
        metrics["narratives"]["split"],
        "",
        metrics["narratives"]["performance"],
        "",
        metrics["narratives"]["per_class"],
        "",
        metrics["narratives"]["robustness"],
        "",
        metrics["narratives"]["ablation"],
        "",
        metrics["narratives"]["calibration"],
        "",
        metrics["narratives"]["importance"],
        "",
        "## Test metrics",
        "",
        "| Model | Accuracy | Macro-F1 | Weighted-F1 | High called low |",
        "| --- | ---: | ---: | ---: | ---: |",
    ]
    for name in ["logistic_regression", "random_forest", "gradient_boosting", "svm_rbf", "soft_voting"]:
        block = metrics["models"][name]["test"]
        lines.append(
            f"| {name} | {_pct(block['accuracy'])} | {_pct(block['macro_f1'])} | {_pct(block['weighted_f1'])} | {block['high_risk_predicted_low']} |"
        )
    lines.append("")
    return "\n".join(lines)


def _add_text(slide, text, left, top, width, height, size=18, bold=False, color=PptColor(0x1C, 0x28, 0x33)):
    box = slide.shapes.add_textbox(PptInches(left), PptInches(top), PptInches(width), PptInches(height))
    frame = box.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = PptPt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.font.name = "Calibri"
    return frame


def build_presentation(bundle: dict, path: Path) -> None:
    metrics = bundle["metrics"]
    group = bundle["group"]
    soft = metrics["models"]["soft_voting"]["test"]
    presentation = Presentation()
    presentation.slide_width = PptInches(13.333)
    presentation.slide_height = PptInches(7.5)
    blank = presentation.slide_layouts[6]

    def slide(title: str, body: str, image: Path | None = None):
        page = presentation.slides.add_slide(blank)
        _add_text(page, title, 0.5, 0.3, 12.3, 0.7, size=28, bold=True, color=PptColor(0x1E, 0x4D, 0x36))
        if image and image.exists():
            _add_text(page, body, 0.5, 1.1, 5.6, 5.6, size=16)
            page.shapes.add_picture(str(image), PptInches(6.4), PptInches(1.4), width=PptInches(6.3))
        else:
            _add_text(page, body, 0.5, 1.2, 12.2, 5.6, size=18)

    names = ", ".join(member["name"] for member in group["members"])
    slide(
        "GreenScope — Lawn Health & Service Planning",
        "An Explainable Ensemble Learning Framework for Plant Health Risk Prediction and Cost-Aware Landscaping Decisions.\n\n"
        f"Group {group['group_number']}\n{names}\n{group['institute']}\n{group['course']}\n{group['submission_date']}\n\n"
        "Independent academic prototype. Not affiliated with The Grounds Guys.",
    )
    slide(
        "Problem",
        "A crew can record a lawn inspection. Deciding the risk, the next action, and a budget is still opaque.\n\n"
        "The course context is the operational problem of a lawn-care business. This software is not that company's product and uses none of its customer records.",
    )
    slide(
        "Questions",
        "Can four different classifiers and a soft vote assess simulated lawn-health risk in a way we can explain?\n\n"
        "Does the vote beat the best single model?\n\n"
        "Which measurements does the model use?\n\n"
        "What happens under a hotter, drier scenario?\n\n"
        "How do we turn the result into actions and an illustrative cost without prescribing chemicals or claiming profit?",
    )
    slide(
        "System",
        "Inspection values pass through each saved pipeline.\n\n"
        "Logistic regression, random forest, gradient boosting, and a calibrated RBF SVM each produce a class and a probability.\n\n"
        "Soft voting averages the probabilities.\n\n"
        "A local explanation, three written screens, actions, and an editable cost follow.\n\n"
        "The research pages use the same saved models and the same test rows.",
    )
    slide(
        "Data, stated plainly",
        "10,000 synthetic inspections and 2,000 robustness rows, both generated here.\n\n"
        "Extension pages informed the meaning of sun, drainage, pH, soil tests, water, mowing, and thatch. They did not supply the rows.\n\n"
        "The risk label is a noisy score of interacting measurements, cut into tertiles on the primary file. It is not a diagnosis.\n\n"
        "No customer identifiers are stored.",
    )
    slide(
        "Why four models",
        "Logistic regression: a linear baseline.\n\n"
        "Random forest: nonlinear splits, depth-limited after a validation pilot showed memorisation.\n\n"
        "Gradient boosting: sequential trees. The tree count and learning rate were chosen on validation only.\n\n"
        "SVM: an RBF kernel with probabilities from training-fold calibration.\n\n"
        "Same target, same rows, same split.",
    )
    slide(
        "Voting",
        "Soft voting: the mean of the four class probabilities. Equal weights. Not tuned on the test set.\n\n"
        "If soft voting does not win, the report says so. The ensemble was pre-specified.",
    )
    slide(
        "Test design",
        metrics["narratives"]["split"] + "\n\nAccuracy, macro-F1, weighted-F1, class-wise scores, confusion matrices, calibration, and a bootstrap of the stored predictions. The bootstrap does not refit the models.",
    )
    slide(
        "Test performance",
        metrics["narratives"]["performance"],
        FIGURE_DIR / "fig_model_comparison.png",
    )
    slide(
        "Where the errors are",
        metrics["narratives"]["per_class"]
        + f"\n\nSoft voting calls {soft['high_risk_predicted_low']} high-risk test rows low risk.",
        FIGURE_DIR / "fig_confusion_soft_voting.png",
    )
    slide(
        "What the model uses",
        metrics["narratives"]["importance"] + "\n\nA local explanation is the change in probability when one measurement is replaced by the fitting-set baseline. It is not a cause.",
        FIGURE_DIR / "fig_feature_importance.png",
    )
    slide(
        "Screens and costs",
        "Water, nutrient, and drainage flags are rules, not extra models.\n\n"
        "Actions ask for a test, an inspection, or a follow-up. They do not state a chemical rate.\n\n"
        "Costs are flat fees plus a rate times lawn area. The defaults are illustrative and editable.\n\n"
        "A later visit can be compared with an earlier one. A new class means the inputs changed, not that a treatment worked.",
    )
    slide(
        "Robustness and ablation",
        metrics["narratives"]["robustness"] + "\n\n" + metrics["narratives"]["ablation"],
        FIGURE_DIR / "fig_robustness.png",
    )
    slide(
        "Limits",
        "Synthetic labels are not field truth.\n\n"
        "Maryland and New Jersey guidance is not an Indian price survey.\n\n"
        "Project nutrient units are not a laboratory method.\n\n"
        "Imputed soil tests are not measurements.\n\n"
        "Calibration and bootstrap intervals describe this benchmark, not customer uncertainty.",
    )
    slide(
        "Conclusion",
        metrics["narratives"]["performance"]
        + "\n\nUse disagreement as a reason to look again. Do not turn the class into a fertiliser rate. "
        "A real labelling study is required before any field-accuracy or business claim.",
    )
    slide(
        "References",
        "University of Maryland Extension, Starting a New Lawn; Maintaining an Established Lawn.\n"
        "Rutgers FS797, FS633, FS102.\n"
        "FAO EcoCrop home page, accessed 22 October 2026. No species records downloaded.\n\n"
        "Full URLs are in DATA_SOURCES.md.",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def copy_supporting_documents() -> None:
    SUBMISSION_DIR.mkdir(parents=True, exist_ok=True)
    mapping = {
        ROOT / "HELP.md": SUBMISSION_DIR / "HELP.md",
        ROOT / "DATA_SOURCES.md": SUBMISSION_DIR / "DATA_SOURCES.md",
        DATA_DIR / "data_dictionary.md": SUBMISSION_DIR / "DATA_DICTIONARY.md",
        DATA_DIR / "source_register.md": ROOT / "DATA_SOURCES.md",
    }
    # Source register is canonical. Copy it to the two DATA_SOURCES paths and copy HELP and the dictionary.
    shutil.copyfile(DATA_DIR / "source_register.md", ROOT / "DATA_SOURCES.md")
    shutil.copyfile(DATA_DIR / "source_register.md", SUBMISSION_DIR / "DATA_SOURCES.md")
    if (ROOT / "HELP.md").exists():
        shutil.copyfile(ROOT / "HELP.md", SUBMISSION_DIR / "HELP.md")
    if (DATA_DIR / "data_dictionary.md").exists():
        shutil.copyfile(DATA_DIR / "data_dictionary.md", SUBMISSION_DIR / "DATA_DICTIONARY.md")
    mapping  # kept so the intended destinations stay visible to a reader of this function


def write_methodology_markdown(bundle: dict) -> None:
    run = bundle["run"]
    text = f"""# Model methodology

This note matches the saved run `{run['created_at_utc']}` and generator `{run['generator_version']}`.

## Target

`overall_lawn_health_risk` with classes Low, Moderate, and High, in that order. All four models and the soft-voting ensemble use this target.

## Split

{bundle['metrics']['narratives']['split']}

## Models

{chr(10).join(f"- {name}: {value}" for name, value in run['hyperparameters'].items())}

## Protocol

{bundle['metrics']['evaluation_protocol']}

## What is not claimed

The test metrics are not field accuracy. The probabilities are not customer confidence. The cost sheet is not a saving.
"""
    (SUBMISSION_DIR / "MODEL_METHODOLOGY.md").write_text(text, encoding="utf-8")


def write_contributions(bundle: dict) -> None:
    group = bundle["group"]
    lines = [
        "# Contributions",
        "",
        "Proposed roles for a five-member group. Replace the names, enrollment numbers, and contribution sentences after the group confirms them. Do not treat a proposed role as evidence that a named person has already done the work.",
        "",
        f"Group number: {group['group_number']}",
        "",
        "| Name | Enrollment | Proposed role | Contribution to confirm | Report location |",
        "| --- | --- | --- | --- | --- |",
    ]
    for member in group["members"]:
        lines.append(
            f"| {member['name']} | {member['enrollment']} | {member['proposed_role']} | {member['contribution']} | {member['report_location']} |"
        )
    lines.append("")
    lines.append("Everyone should review the interpretation, the report, and the presentation.")
    lines.append("")
    (SUBMISSION_DIR / "CONTRIBUTIONS_TEMPLATE.md").write_text("\n".join(lines), encoding="utf-8")
    (ROOT / "GROUP_INFO_TEMPLATE.md").write_text(
        "\n".join(
            [
                "# Group information",
                "",
                "Edit `group_info.json`, then run `python -m src.build_documents`.",
                "The report's first page is rebuilt from that file.",
                "",
                "Required before submission:",
                "",
                "- Group number.",
                "- Five names and enrollment numbers. Do not invent them.",
                "- A confirmed one- or two-sentence contribution for each member.",
                "- Printed page numbers checked against the paginated report. The template gives section locations because pagination depends on the word processor.",
                "- Semester, instructor, and submission date if the course requires them.",
                "",
                "Proposed roles, to be confirmed rather than assumed:",
                "",
                "1. Data sources, dictionary, generator, provenance, and validation.",
                "2. Logistic regression, random forest, and preprocessing.",
                "3. Gradient boosting, SVM, soft voting, and ablation.",
                "4. Metrics, confusion matrices, robustness, and explanations.",
                "5. Streamlit application, demonstration, recommendations, costs, and help.",
                "",
            ]
        ),
        encoding="utf-8",
    )


def main() -> None:
    bundle = _load()
    report = build_report(bundle)
    markdown = report.markdown()
    (ROOT / "PROJECT_REPORT.md").write_text(markdown, encoding="utf-8")
    (SUBMISSION_DIR / "PROJECT_REPORT.md").write_text(
        markdown.replace("](reports/", "](../reports/"),
        encoding="utf-8",
    )
    report.docx(SUBMISSION_DIR / "PROJECT_REPORT.docx")
    (REPORT_DIR / "experiment_analysis.md").write_text(build_experiment_markdown(bundle), encoding="utf-8")
    build_presentation(bundle, SUBMISSION_DIR / "PROJECT_PRESENTATION.pptx")
    copy_supporting_documents()
    write_methodology_markdown(bundle)
    write_contributions(bundle)
    print(f"Wrote report, presentation, and submission files to {SUBMISSION_DIR}")


if __name__ == "__main__":
    main()
